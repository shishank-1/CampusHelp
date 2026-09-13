"""
CampusHelp - Phase 1: Data Ingestion
--------------------------------------
Reads PDF and DOCX files from a folder, extracts raw text, and splits
that text into overlapping chunks ready for embedding in a later phase.

Usage:
    python ingest.py                # processes everything in data/raw_docs/
"""

import json
from pathlib import Path

from pypdf import PdfReader
from docx import Document

# ---------- Config ----------
RAW_DOCS_DIR = Path(__file__).parent.parent / "data" / "raw_docs"
PROCESSED_DIR = Path(__file__).parent.parent / "data" / "processed"

CHUNK_SIZE = 800        # characters per chunk
CHUNK_OVERLAP = 100     # characters of overlap between consecutive chunks


# ---------- Text extraction ----------
def extract_text_from_pdf(file_path: Path) -> str:
    """Extract all text from a PDF file, page by page."""
    reader = PdfReader(str(file_path))
    pages_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)
    return "\n".join(pages_text)


def extract_text_from_docx(file_path: Path) -> str:
    """Extract paragraph and table text from a DOCX file."""
    doc = Document(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))

    return "\n".join(paragraphs + table_rows)


def extract_text(file_path: Path) -> str:
    """Dispatch to the right extractor based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    elif suffix == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix} ({file_path.name})")


# ---------- Cleaning ----------
def clean_text(text: str) -> str:
    """Basic cleanup: collapse excessive whitespace/blank lines."""
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]  # drop empty lines
    return "\n".join(lines)


# ---------- Chunking ----------
def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """
    Split text into overlapping chunks of roughly `chunk_size` characters.
    Overlap helps preserve context across chunk boundaries.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size")

    chunks = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = start + chunk_size
        chunk = text[start:end]
        chunks.append(chunk.strip())

        if end >= text_length:
            break

        # Move start forward, leaving `overlap` characters of repeat
        start = end - overlap

    return [c for c in chunks if c]  # drop any empty chunks


# ---------- Main pipeline ----------
def process_document(file_path: Path) -> dict:
    """Extract, clean, and chunk a single document. Returns a record dict."""
    print(f"Processing: {file_path.name}")

    raw_text = extract_text(file_path)
    cleaned = clean_text(raw_text)
    chunks = chunk_text(cleaned)

    record = {
        "source_file": file_path.name,
        "num_characters": len(cleaned),
        "num_chunks": len(chunks),
        "chunks": [
            {"chunk_id": f"{file_path.stem}_{i}", "text": chunk, "source": file_path.name}
            for i, chunk in enumerate(chunks)
        ],
    }
    return record


def run_ingestion():
    """Process every PDF/DOCX file in data/raw_docs/ and save results as JSON."""
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

    if not RAW_DOCS_DIR.exists():
        print(f"Raw documents folder not found: {RAW_DOCS_DIR}")
        print("Create the folder, add PDF/DOCX files, and run this script again.")
        return

    supported_files = [
        f for f in RAW_DOCS_DIR.iterdir()
        if f.suffix.lower() in (".pdf", ".docx")
    ]
    supported_files.sort(key=lambda path: path.name.lower())

    if not supported_files:
        print(f"No PDF/DOCX files found in {RAW_DOCS_DIR}")
        print("Add some documents there and run this script again.")
        return

    all_records = []
    for file_path in supported_files:
        try:
            record = process_document(file_path)
            all_records.append(record)
            print(f"  -> {record['num_chunks']} chunks created ({record['num_characters']} chars)")
        except Exception as e:
            print(f"  !! Failed to process {file_path.name}: {e}")

    output_path = PROCESSED_DIR / "chunks.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_records, f, indent=2, ensure_ascii=False)

    total_chunks = sum(r["num_chunks"] for r in all_records)
    print(f"\nDone. {len(all_records)} document(s) processed, {total_chunks} total chunks.")
    print(f"Saved to: {output_path}")


if __name__ == "__main__":
    run_ingestion()